import os, sys, json, ast, subprocess, importlib

ROOT=os.getcwd()
FAIL=[]
PASS=[]

def ok(x):
    PASS.append(x)
    print("[PASS]",x)

def fail(x):
    FAIL.append(x)
    print("[FAIL]",x)

# ---------- CORE ----------
required=[
"anviqo_product.py","anviqo_api.py","anviqo_web.py",
"anvi_knowledge_layer.py","pci_spares.py","pci_conversation.py",
"pci_live_simulator.py","anvi_field_report.py","anvi_voice.py",
"plant_memory.py","event_correlation.py","event_timeline.py",
"plant_health.py","plant_brain_reasoning.py"
]
for f in required:
    if os.path.isfile(f): ok("Core file: "+f)
    else: fail("Missing core file: "+f)

# ---------- COMPILE ACTIVE PYTHON ----------
for f in required:
    if os.path.isfile(f):
        try:
            ast.parse(open(f,encoding="utf-8").read(),filename=f)
            ok("Compile: "+f)
        except Exception as e:
            fail("Compile: "+f+" -> "+str(e))

# ---------- PCI ----------
pci="database/pci/pci_instrument_database.json"
try:
    d=json.load(open(pci,encoding="utf-8"))
    assert d.get("version")=="PCI-1.0"
    assert d.get("status")=="REFERENCE_ONLY"
    assert d.get("control_mode")=="READ_ONLY"
    assert d.get("plc_write") is False
    assert d.get("scada_control") is False
    assert d.get("human_decision_required") is True
    assert d.get("record_count")==1064
    ok("PCI database and safety contract")
except Exception as e:
    fail("PCI database: "+str(e))

# ---------- SPARES ----------
try:
    import openpyxl
    wb=openpyxl.load_workbook("database/spares/critical_spares.xlsx")
    assert "PT" in wb.sheetnames
    ws=wb["PT"]
    found=None
    for r in range(1,ws.max_row+1):
        if str(ws.cell(r,1).value).strip().upper() in ("PT-303","PT303"):
            found=r
            break
    if found is None:
        # search entire row
        for r in range(1,ws.max_row+1):
            vals=[str(ws.cell(r,c).value).upper() for c in range(1,ws.max_column+1)]
            if any("PT-303" in v or "PT303" in v for v in vals):
                found=r; break
    assert found is not None
    ok("Critical spare workbook/PT-303")
except Exception as e:
    fail("Critical spare workbook: "+str(e))

# ---------- RUNTIME SAFETY ----------
safety_files=[
"anviqo_api.py","anvi_knowledge_layer.py","anvi_field_report.py",
"pci_spares.py","pci_conversation.py","pci_live_simulator.py","anvi_voice.py"
]
for f in safety_files:
    try:
        text=open(f,encoding="utf-8").read().lower()
        forbidden=[
            "plc_write = true","plc_write=true",
            "scada_control = true","scada_control=true",
            "real_plc_write_enabled = true","real_plc_write_enabled=true",
            "real_scada_control_enabled = true","real_scada_control_enabled=true"
        ]
        bad=[x for x in forbidden if x in text]
        if bad: fail(f+" contains unsafe enable")
        else: ok(f+" no PLC/SCADA enable")
    except Exception as e:
        fail(f+" safety check: "+str(e))

# ---------- SPARE EXECUTION CONTRACT ----------
try:
    import pci_spares
    assert hasattr(pci_spares,"execute_spare_mutation_v18")
    src=open("pci_spares.py",encoding="utf-8").read()
    assert "execute_spare_mutation_v18" in src
    assert "negative" in src.lower()
    ok("V1.8 direct spare mutation engine")
except Exception as e:
    fail("V1.8 spare engine: "+str(e))

# ---------- MEMORY ----------
try:
    pm=json.load(open("database/plant_memory/plant_memory.json",encoding="utf-8"))
    if isinstance(pm,dict):
        records=pm.get("memories",pm.get("records",pm.get("data",[])))
    else:
        records=pm
    assert isinstance(records,list)
    ok("Plant Memory JSON structure")
except Exception as e:
    fail("Plant Memory: "+str(e))

# ---------- ACTIVE SECRET CHECK ----------
bad_secret=False
for f in ["anviqo_api.py","anviqo_web.py"]:
    if os.path.isfile(f):
        t=open(f,encoding="utf-8").read()
        if "app.secret_key = os.environ.get" in t or "app.config['SECRET_KEY'] = os.environ.get" in t or 'app.config["SECRET_KEY"] = os.environ.get' in t:
            ok(f+" environment-based secret")
        elif "secret_key = os.environ.get" in t.lower():
            ok(f+" environment-based secret")
        else:
            fail(f+" environment secret configuration not detected")
            bad_secret=True

# ---------- AUTH ----------
try:
    t=open("anviqo_api.py",encoding="utf-8").read()
    assert "ADMIN_USER" in t and "ADMIN_PASSWORD" in t
    assert "login_required" in t
    ok("Authentication configuration")
except Exception as e:
    fail("Authentication: "+str(e))

# ---------- ROUTES ----------
try:
    t=open("anviqo_api.py",encoding="utf-8").read()
    routes=[
        "/api/status","/api/pci","/api/pci/live","/api/ask",
        "/api/equipment/<tag>","/api/maintenance",
        "/api/management","/api/plant_snapshot","/api/field_report"
    ]
    for r in routes:
        if r not in t: fail("Missing route "+r)
    else:
        ok("Required API routes")
except Exception as e:
    fail("Routes: "+str(e))

# ---------- GIT ----------
try:
    r=subprocess.run(["git","status","--porcelain"],capture_output=True,text=True)
    print("===== RELEASE CHANGES =====")
    print(r.stdout)
except Exception as e:
    fail("Git status: "+str(e))

print("================================================")
print("ANVIQO RELEASE GATE V2")
print("PASS :",len(PASS))
print("FAIL :",len(FAIL))
print("================================================")

if FAIL:
    print("RELEASE GATE: FAIL")
    for x in FAIL:
        print(" -",x)
    sys.exit(1)

print("RELEASE GATE: PASS")
