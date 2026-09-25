from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell

import pci_spare_direct_excel as engine


def test_quantity_merge_normalization_keeps_owner_only():
    wb = Workbook()
    ws = wb.active
    ws.title = "PT"
    ws.append(["Tag No", "Qty Avbl"])
    ws.append(["PT-519", 4])
    ws.append(["PT-520", None])
    ws.merge_cells("B2:B3")

    changed = engine._normalize_quantity_merges(wb)

    assert changed is True
    assert not ws.merged_cells.ranges
    assert ws["B2"].value == 4
    assert ws["B3"].value is None
    assert not isinstance(ws["B2"], MergedCell)
    assert not isinstance(ws["B3"], MergedCell)


def test_normalization_never_duplicates_shared_quantity():
    wb = Workbook()
    ws = wb.active
    ws.title = "PT"
    ws.append(["Tag No", "Qty Avbl"])
    ws.append(["PT-519", 4])
    ws.append(["PT-520", None])
    ws.append(["PT-521", None])
    ws.merge_cells("B2:B4")

    engine._normalize_quantity_merges(wb)

    assert [ws.cell(r, 2).value for r in range(2, 5)] == [4, None, None]
