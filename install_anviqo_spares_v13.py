from pathlib import Path
from datetime import datetime, timezone
import shutil
import json
import re
import openpyxl

ROOT = Path(__file__).resolve().parent
CONV = ROOT / "pci_conversation.py"
SPARE_DIR = ROOT / "database" / "spares"

NORMAL_XLSX = SPARE_DIR / "normal_spares.xlsx"
NORMAL_AUDIT = SPARE_DIR / "normal_spare_transactions.json"
UNIVERSAL = ROOT / "pci_spare_universal_v13.py"

STAMP = datetime.now().strftime("%Y%m%d-%H%M%S")


# ============================================================
# BACKUP
# ============================================================

if CONV.exists():
    backup = CONV.with_name(
        f"pci_conversation.py.backup-before-spares-v13-{STAMP}"
    )
    shutil.copy2(CONV, backup)
    print("BACKUP:", backup)


SPARE_DIR.mkdir(parents=True, exist_ok=True)


# ============================================================
# NORMAL SPARES WORKBOOK
# ============================================================

if not NORMAL_XLSX.exists():

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Normal Spares"

    headers = [
        "ID",
        "ITEM",
        "DESCRIPTION",
        "PART/MODEL",
        "CATEGORY",
        "LOCATION",
        "QUANTITY",
        "MINIMUM STOCK",
        "UNIT",
        "REMARKS"
    ]

    ws.append(headers)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(NORMAL_XLSX)

    print("CREATED:", NORMAL_XLSX)

else:
    print("EXISTS:", NORMAL_XLSX)


# ============================================================
# NORMAL AUDIT
# ============================================================

if not NORMAL_AUDIT.exists():
    NORMAL_AUDIT.write_text(
        json.dumps({"transactions": []}, indent=2),
        encoding="utf-8"
    )
    print("CREATED:", NORMAL_AUDIT)


# ============================================================
# UNIVERSAL ENGINE
# ============================================================

ENGINE = r'''
from pathlib import Path
from datetime import datetime, timezone
import json
import re
import openpyxl

ROOT = Path(__file__).resolve().parent

CRITICAL_XLSX = ROOT / "database" / "spares" / "critical_spares.xlsx"
CRITICAL_AUDIT = ROOT / "database" / "spares" / "critical_spare_transactions.json"

NORMAL_XLSX = ROOT / "database" / "spares" / "normal_spares.xlsx"
NORMAL_AUDIT = ROOT / "database" / "spares" / "normal_spare_transactions.json"


def norm(v):
    return re.sub(r"[^a-z0-9]+", " ", str(v or "").lower()).strip()


def number(v):
    try:
        return float(v)
    except Exception:
        return 0.0


def fmt_qty(v):
    v = number(v)
    return str(int(v)) if v.is_integer() else str(v)


def audit_file(path):
    if not path.exists():
        return {"transactions": []}

    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {"transactions": []}


def save_audit(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(
        json.dumps(data, indent=2, ensure_ascii=False),
        encoding="utf-8"
    )
    tmp.replace(path)


def write_audit(path, action, tag, qty, before, after, source):
    data = audit_file(path)

    txn = (
        "SPARE-V13-"
        + datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S%f")[:-3]
    )

    data.setdefault("transactions", []).append({
        "transaction_id": txn,
        "action": action,
        "tag": tag,
        "quantity": qty,
        "quantity_before": before,
        "quantity_after": after,
        "status": "APPLIED",
        "source": source,
        "excel_authoritative": True,
        "plc_write": False,
        "scada_control": False,
        "timestamp_utc": datetime.now(timezone.utc).isoformat()
    })

    save_audit(path, data)
    return txn


def headers(ws):
    result = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                result[norm(cell.value)] = cell.column
        if result:
            break
    return result


def find_header(ws, names):
    wanted = {norm(x) for x in names}

    for row in ws.iter_rows():
        for cell in row:
            if norm(cell.value) in wanted:
                return cell.column

    return None


def find_tag(ws, tag):
    target = norm(tag)

    # Exact match first
    for row in ws.iter_rows():
        for cell in row:
            if norm(cell.value) == target:
                return cell.row

    # Then substring
    for row in ws.iter_rows():
        for cell in row:
            value = norm(cell.value)
            if target and target in value:
                return cell.row

    return None


def critical_records(query=None):
    if not CRITICAL_XLSX.exists():
        return []

    wb = openpyxl.load_workbook(CRITICAL_XLSX, data_only=False)
    records = []

    for ws in wb.worksheets:

        tag_col = find_header(ws, [
            "TAG NO", "TAG", "TAGNO"
        ])

        item_col = find_header(ws, [
            "SHORT DESCRIPTION",
            "SHORT DESCRIPTION",
            "INSTRUMENT",
            "DESCRIPTION",
            "Item"
        ])

        qty_col = find_header(ws, [
            "Qty avbl",
            "Qty available",
            "Qty Avbl",
            "AVAILABLE",
            "Qty"
        ])

        req_col = find_header(ws, [
            "Qty Req",
            "Qty Req.",
            "Qty Required"
        ])

        indent_col = find_header(ws, [
            "Spare to be indent"
        ])

        if not qty_col:
            continue

        for row in range(1, ws.max_row + 1):

            values = [
                ws.cell(row, c).value
                for c in range(1, ws.max_column + 1)
            ]

            text = " ".join(norm(v) for v in values if v is not None)

            if query and norm(query) not in text:
                continue

            tag = (
                ws.cell(row, tag_col).value
                if tag_col else None
            )

            item = (
                ws.cell(row, item_col).value
                if item_col else None
            )

            qty = number(ws.cell(row, qty_col).value)

            req = (
                number(ws.cell(row, req_col).value)
                if req_col else 0
            )

            indent = (
                ws.cell(row, indent_col).value
                if indent_col else None
            )

            if tag is None and item is None:
                continue

            records.append({
                "source": "critical",
                "sheet": ws.title,
                "row": row,
                "tag": str(tag or ""),
                "item": str(item or ""),
                "quantity": qty,
                "required": req,
                "indent": indent
            })

    return records


def normal_records(query=None):
    if not NORMAL_XLSX.exists():
        return []

    wb = openpyxl.load_workbook(NORMAL_XLSX, data_only=False)
    ws = wb["Normal Spares"]

    records = []

    for row in range(2, ws.max_row + 1):

        item = ws.cell(row, 2).value
        desc = ws.cell(row, 3).value
        part = ws.cell(row, 4).value
        category = ws.cell(row, 5).value
        location = ws.cell(row, 6).value
        qty = number(ws.cell(row, 7).value)
        minimum = number(ws.cell(row, 8).value)
        unit = ws.cell(row, 9).value

        text = " ".join(
            norm(x)
            for x in [item, desc, part, category, location]
            if x is not None
        )

        if query and norm(query) not in text:
            continue

        if not item and not desc and not part:
            continue

        records.append({
            "source": "normal",
            "row": row,
            "tag": str(item or ""),
            "item": str(item or ""),
            "description": str(desc or ""),
            "part": str(part or ""),
            "category": str(category or ""),
            "location": str(location or ""),
            "quantity": qty,
            "required": minimum,
            "unit": str(unit or "")
        })

    return records


def all_records(query=None):
    return critical_records(query) + normal_records(query)


def stock_status(qty, required=0):
    if qty <= 0:
        return "ZERO STOCK"

    if required > 0 and qty < required:
        return "BELOW REQUIRED"

    if required > 0 and qty >= required:
        return "SUFFICIENT STOCK"

    return "AVAILABLE"


def search(query):
    records = all_records(query)

    if not records:
        return "No spare records found for: " + str(query)

    lines = [
        f"ANVI found {len(records)} spare record(s)."
    ]

    for r in records:

        status = stock_status(
            r["quantity"],
            r.get("required", 0)
        )

        identifier = r.get("tag") or r.get("item")

        lines.append(
            f"{identifier} — "
            f"Qty: {fmt_qty(r['quantity'])}; "
            f"Status: {status}; "
            f"Source: {r['source']}"
        )

    return "\n".join(lines)


def find_exact(tag):
    target = norm(tag)

    for r in all_records():
        if norm(r.get("tag")) == target:
            return r

    return None


def critical_update(tag, quantity, action):
    from pci_spare_direct_excel import update_spare

    result = update_spare(tag, quantity, action)

    action_text = (
        "Added" if action == "ADD"
        else "Used/removed"
    )

    return (
        f"{tag.upper()} spare quantity updated directly in Excel.\n"
        f"Previous quantity: {fmt_qty(result['before'])}\n"
        f"{action_text}: {fmt_qty(result['quantity'])}\n"
        f"Current quantity: {fmt_qty(result['after'])}\n"
        f"Direct Excel update: APPLIED\n"
        f"Audit ID: {result['transaction_id']}"
    )


def normal_add(item, quantity, description="",
               part="", category="", location="",
               minimum=0, unit="Nos", remarks=""):

    wb = openpyxl.load_workbook(NORMAL_XLSX)
    ws = wb["Normal Spares"]

    existing = None

    for row in range(2, ws.max_row + 1):
        if norm(ws.cell(row, 2).value) == norm(item):
            existing = row
            break

    if existing:
        row = existing
        before = number(ws.cell(row, 7).value)
        after = before + float(quantity)
        ws.cell(row, 7).value = int(after) if after.is_integer() else after
        action = "ADD"
    else:
        row = ws.max_row + 1
        before = 0.0
        after = float(quantity)

        ws.cell(row, 1).value = f"NS-{row-1:04d}"
        ws.cell(row, 2).value = item
        ws.cell(row, 3).value = description
        ws.cell(row, 4).value = part
        ws.cell(row, 5).value = category
        ws.cell(row, 6).value = location
        ws.cell(row, 7).value = int(after) if after.is_integer() else after
        ws.cell(row, 8).value = minimum
        ws.cell(row, 9).value = unit
        ws.cell(row, 10).value = remarks

        action = "ADD_NEW"

    tmp = NORMAL_XLSX.with_suffix(".v13.tmp.xlsx")
    wb.save(tmp)
    tmp.replace(NORMAL_XLSX)

    txn = write_audit(
        NORMAL_AUDIT,
        action,
        item,
        float(quantity),
        before,
        after,
        "normal_spares.xlsx"
    )

    return (
        f"Normal spare '{item}' updated directly in Excel.\n"
        f"Previous quantity: {fmt_qty(before)}\n"
        f"Added: {fmt_qty(quantity)}\n"
        f"Current quantity: {fmt_qty(after)}\n"
        f"Direct Excel update: APPLIED\n"
        f"Audit ID: {txn}"
    )


def normal_remove(item, quantity):

    record = find_exact(item)

    if not record or record["source"] != "normal":
        raise ValueError(
            f"Normal spare {item} was not found."
        )

    wb = openpyxl.load_workbook(NORMAL_XLSX)
    ws = wb["Normal Spares"]

    row = record["row"]
    before = number(ws.cell(row, 7).value)
    quantity = float(quantity)

    if quantity > before:
        raise ValueError(
            f"Cannot remove {fmt_qty(quantity)} {item}. "
            f"Only {fmt_qty(before)} available."
        )

    after = before - quantity

    ws.cell(row, 7).value = (
        int(after) if after.is_integer() else after
    )

    tmp = NORMAL_XLSX.with_suffix(".v13.tmp.xlsx")
    wb.save(tmp)
    tmp.replace(NORMAL_XLSX)

    txn = write_audit(
        NORMAL_AUDIT,
        "REMOVE",
        item,
        quantity,
        before,
        after,
        "normal_spares.xlsx"
    )

    return (
        f"Normal spare '{item}' updated directly in Excel.\n"
        f"Previous quantity: {fmt_qty(before)}\n"
        f"Used/removed: {fmt_qty(quantity)}\n"
        f"Current quantity: {fmt_qty(after)}\n"
        f"Direct Excel update: APPLIED\n"
        f"Audit ID: {txn}"
    )


def indent_report():

    records = all_records()
    short = []

    for r in records:

        qty = r["quantity"]
        required = r.get("required", 0)

        if qty <= 0 or (required > 0 and qty < required):
            short.append(r)

    if not short:
        return "ANVI: No spare currently identified as short."

    lines = [
        f"ANVI identified {len(short)} spare record(s) requiring attention:"
    ]

    for r in short:
        identifier = r.get("tag") or r.get("item")

        lines.append(
            f"{identifier} — "
            f"Available: {fmt_qty(r['quantity'])}; "
            f"Required: {fmt_qty(r.get('required', 0))}; "
            f"Status: {stock_status(r['quantity'], r.get('required', 0))}"
        )

    return "\n".join(lines)


def equipment_search(equipment):

    q = norm(equipment)

    records = []

    for r in all_records():
        text = norm(
            " ".join(
                str(r.get(k, ""))
                for k in [
                    "tag",
                    "item",
                    "description",
                    "part",
                    "category"
                ]
            )
        )

        if q in text:
            records.append(r)

    if not records:
        return f"No spare relationship records found for {equipment}."

    lines = [
        f"ANVI found {len(records)} spare record(s) related to {equipment}:"
    ]

    for r in records:
        identifier = r.get("tag") or r.get("item")
        lines.append(
            f"{identifier} — Qty: {fmt_qty(r['quantity'])}; "
            f"Status: {stock_status(r['quantity'], r.get('required', 0))}"
        )

    return "\n".join(lines)


def router(question):

    q = str(question or "").strip()
    n = norm(q)

    # --------------------------------------------------------
    # INDENT
    # --------------------------------------------------------

    if (
        "show spares to indent" in n
        or "spares to indent" in n
        or "which spares are short" in n
        or "which spares are shortage" in n
        or "spares are short" in n
    ):
        return indent_report()

    # --------------------------------------------------------
    # UNIVERSAL SEARCH
    # --------------------------------------------------------

    if (
        "show all spares" in n
        or "search all spares" in n
        or "universal spare search" in n
    ):
        return search("")

    # --------------------------------------------------------
    # NATURAL QUANTITY UPDATE
    # --------------------------------------------------------

    number_match = re.search(
        r"\b(\d+(?:\.\d+)?)\s*(?:nos?|number|numbers|pcs?|pieces?)?\b",
        n
    )

    qty = float(number_match.group(1)) if number_match else None

    # Extract common tag pattern
    tag_match = re.search(
        r"\b(?:pt|ft|lt|tt|te|le|fe|mcv|sov|fsv|pcv|wt|sv)"
        r"[\s\-_]?\d+[a-z0-9/\-_]*",
        q,
        re.I
    )

    tag = tag_match.group(0).upper() if tag_match else None

    if tag and qty is not None:

        is_add = any(x in n for x in [
            "add",
            "received",
            "receive",
            "new spare",
            "new equipment"
        ])

        is_remove = any(x in n for x in [
            "used",
            "i used",
            "have used",
            "remove",
            "removed",
            "consume",
            "consumed",
            "issued",
            "issue"
        ])

        if is_add or is_remove:

            record = find_exact(tag)

            # Existing critical spare
            if record and record["source"] == "critical":

                if is_add:
                    return critical_update(tag, qty, "ADD")

                return critical_update(tag, qty, "REMOVE")

            # Existing normal spare
            if record and record["source"] == "normal":

                if is_add:
                    return normal_add(tag, qty)

                return normal_remove(tag, qty)

    # --------------------------------------------------------
    # NEW NORMAL SPARE
    # --------------------------------------------------------

    m = re.search(
        r"add\s+(?:a\s+)?new\s+normal\s+spare\s+(.+?)"
        r"(?:\s+quantity\s+(\d+(?:\.\d+)?))?$",
        q,
        re.I
    )

    if m:
        item = m.group(1).strip(" ,.")
        quantity = float(m.group(2) or 0)

        if quantity <= 0:
            return (
                "Please specify the quantity, for example: "
                "ANVI, add new normal spare bearing 6205 quantity 5"
            )

        return normal_add(item, quantity)

    # --------------------------------------------------------
    # EQUIPMENT / CATEGORY
    # --------------------------------------------------------

    m = re.search(
        r"(?:spares?\s+for|spares?\s+available\s+for)\s+(.+)$",
        n
    )

    if m:
        return equipment_search(m.group(1).strip())

    # --------------------------------------------------------
    # CURRENT QUANTITY
    # --------------------------------------------------------

    if tag and any(x in n for x in [
        "current spare quantity",
        "current quantity",
        "how many",
        "quantity of",
        "available"
    ]):

        record = find_exact(tag)

        if record:
            return (
                f"{tag} — Available: "
                f"{fmt_qty(record['quantity'])}; "
                f"Status: "
                f"{stock_status(record['quantity'], record.get('required', 0))}; "
                f"Source: {record['source']}"
            )

    return None
'''

UNIVERSAL.write_text(ENGINE, encoding="utf-8")
print("CREATED:", UNIVERSAL)


# ============================================================
# PATCH CONVERSATION ROUTER
# ============================================================

if not CONV.exists():
    raise SystemExit("ERROR: pci_conversation.py not found")

s = CONV.read_text(encoding="utf-8")

if "pci_spare_universal_v13" not in s:

    import_line = (
        "\n# ANVIQO SPARE INTELLIGENCE V1.3\n"
        "from pci_spare_universal_v13 import router as spare_v13_router\n"
    )

    # Insert import after existing imports
    lines = s.splitlines()
    insert_at = 0

    for i, line in enumerate(lines):
        if line.startswith("import ") or line.startswith("from "):
            insert_at = i + 1

    lines.insert(insert_at, import_line.strip())

    s = "\n".join(lines) + "\n"

    # Find answer function
    marker = None

    for candidate in [
        "def answer(question):",
        "def answer(q):",
        "def answer(query):"
    ]:
        if candidate in s:
            marker = candidate
            break

    if marker is None:
        print("WARNING: answer() marker not found.")
        print("Universal engine created but conversation router was NOT modified.")
    else:

        replacement = (
            marker
            + "\n"
            + "    # V1.3 UNIVERSAL SPARE ROUTER\n"
            + "    try:\n"
            + "        _spare_v13 = spare_v13_router(question)\n"
            + "        if _spare_v13:\n"
            + "            return {\"answer\": _spare_v13}\n"
            + "    except Exception as _spare_v13_error:\n"
            + "        # Preserve existing ANVI behavior if V1.3 cannot handle a query.\n"
            + "        pass\n"
        )

        s = s.replace(marker, replacement, 1)

        CONV.write_text(s, encoding="utf-8")

        print("PATCHED: pci_conversation.py")
        print("V1.3 universal spare router connected.")


# ============================================================
# COMPILE CHECK
# ============================================================

import py_compile

for f in [
    UNIVERSAL,
    CONV
]:
    py_compile.compile(
        str(f),
        doraise=True
    )
    print("SYNTAX PASS:", f.name)


print()
print("=" * 60)
print(" ANVIQO SPARE INTELLIGENCE V1.3 INSTALL COMPLETE")
print("=" * 60)
print("Critical Excel :", CRITICAL_XLSX)
print("Normal Excel   :", NORMAL_XLSX)
print("Universal      :", UNIVERSAL)
print("Audit critical :", CRITICAL_AUDIT)
print("Audit normal   :", NORMAL_AUDIT)
print()
print("Existing V1.2 critical-spare engine preserved.")
print("PLC write      : FALSE")
print("SCADA control  : FALSE")
print("=" * 60)
