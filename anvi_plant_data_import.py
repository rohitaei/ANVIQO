"""ANVIQO clean plant-data import path.
Direct synchronous import; no ingestion job, queue, worker or polling.
"""
from __future__ import annotations
import csv, hashlib, io, json, re
from dataclasses import asdict, is_dataclass
from datetime import datetime, timezone
import zipfile
import anvi_tenant_store as store
from universal_onboarding import normalize_record, SAFETY
VERSION = "ANVIQO-DIRECT-PLANT-DATA-V1.4.5"

def _now(): return datetime.now(timezone.utc).isoformat()
def _p(): return store._placeholder()
def _nonempty(v): return v is not None and str(v).strip() != ""
def _key(v): return re.sub(r"[^a-z0-9]+","",str(v or "").strip().lower())
def _header_score(row):
    known={"tag","tagname","plctag","plctagname","externalid","assetid","area","service","description","iotype","plcaddress","panel","tb","jb","recordtype","entitytype"}
    return sum(1 for v in row if _key(v) in known)

def _matrix_records(sheet_name, rows):
    matrix=[tuple(row) for row in rows if any(_nonempty(v) for v in row)]
    if not matrix: return
    header_index=max(range(min(30,len(matrix))),key=lambda i:_header_score(matrix[i]))
    first=matrix[header_index]; headers=[]; seen={}
    for i,v in enumerate(first):
        h=str(v).strip() if _nonempty(v) else f"column_{i+1}"; k=_key(h); seen[k]=seen.get(k,0)+1
        headers.append(h if seen[k]==1 else f"{h}_{seen[k]}")
    for idx,row in enumerate(matrix[header_index+1:],start=header_index+2):
        if not any(_nonempty(v) for v in row): continue
        rec={headers[i]:row[i] for i in range(min(len(headers),len(row)))}
        rec={str(k).strip():v for k,v in rec.items() if str(k).strip()}
        rec["external_id"]=str(rec.get("external_id") or rec.get("id") or rec.get("tag") or rec.get("asset_id") or rec.get("TAG") or hashlib.sha256((sheet_name+str(idx)+json.dumps(rec,sort_keys=True,default=str)).encode()).hexdigest()[:20])
        rec["source"]=sheet_name; rec["record_type"]=str(rec.get("record_type") or rec.get("entity_type") or rec.get("type") or "ASSET")
        yield rec

def _xlsx_records(raw):
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(bytes(raw)),read_only=True,data_only=True)
    try:
        for ws in wb.worksheets:
            yield from _matrix_records(ws.title, ws.iter_rows(values_only=True))
    finally: wb.close()

def _xls_records(raw):
    import xlrd
    wb=xlrd.open_workbook(file_contents=bytes(raw),on_demand=True)
    try:
        for ws in wb.sheets():
            rows=(ws.row_values(i) for i in range(ws.nrows))
            yield from _matrix_records(ws.name, rows)
    finally:
        wb.release_resources()

def _looks_like_xlsx(raw):
    return bytes(raw).startswith(b"PK\x03\x04")
def _looks_like_xls(raw):
    return bytes(raw).startswith(b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1")

def _text(raw,ext):
    if ext in {"txt","log","csv"}: return bytes(raw).decode("utf-8-sig",errors="replace")
    if ext=="docx":
        with zipfile.ZipFile(io.BytesIO(bytes(raw))) as z: xml=z.read("word/document.xml").decode("utf-8",errors="replace")
        return "\n".join(re.sub(r"<[^>]+>"," ",x).strip() for x in re.findall(r"<w:t[^>]*>(.*?)</w:t>",xml) if x.strip())
    if ext=="pdf":
        try:
            from pypdf import PdfReader
            return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(bytes(raw))).pages)
        except Exception: return ""
    return ""

def _records(filename,raw):
    ext=filename.lower().rsplit(".",1)[-1] if "." in filename else ""
    # Prefer the filename extension, but fall back to the actual binary signature.
    # This handles legacy XLS files that were uploaded/stored with an XLSX name.
    if ext=="xls" or (not ext or _looks_like_xls(raw)):
        try: yield from _xls_records(raw); return
        except Exception:
            if ext=="xls": raise
    if ext=="xlsx" or _looks_like_xlsx(raw):
        try: yield from _xlsx_records(raw); return
        except Exception:
            if ext=="xlsx": raise
    if ext=="csv":
        reader=csv.DictReader(io.StringIO(bytes(raw).decode("utf-8-sig",errors="replace")))
        for i,row in enumerate(reader):
            row=dict(row); row["external_id"]=str(row.get("external_id") or row.get("id") or row.get("tag") or row.get("asset_id") or hashlib.sha256((filename+str(i)+json.dumps(row,sort_keys=True,default=str)).encode()).hexdigest()[:20]); row["source"]=filename; row["record_type"]=str(row.get("record_type") or row.get("entity_type") or row.get("type") or "ASSET"); yield row
        return
    if ext=="json":
        data=json.loads(bytes(raw).decode("utf-8",errors="replace")); data=(data.get("records") or data.get("data") or [data]) if isinstance(data,dict) else data
        for i,row in enumerate(data if isinstance(data,list) else []):
            if not isinstance(row,dict): continue
            row=dict(row); row["external_id"]=str(row.get("external_id") or row.get("id") or row.get("tag") or row.get("asset_id") or hashlib.sha256((filename+str(i)+json.dumps(row,sort_keys=True,default=str)).encode()).hexdigest()[:20]); row["source"]=filename; row["record_type"]=str(row.get("record_type") or row.get("entity_type") or row.get("type") or "ASSET"); yield row
        return
    text=_text(raw,ext)
    if text.strip():
        for i in range(0,len(text),6000): yield {"external_id":hashlib.sha256((filename+str(i)).encode()).hexdigest()[:20],"name":filename,"record_type":"DOCUMENT","source":filename,"metadata":{"content":text[i:i+6000]}}

def _knowledge_schema():
    with store._connect() as conn:
        cur=conn.cursor(); p=_p()
        if store._is_sqlite():
            cur.execute("CREATE TABLE IF NOT EXISTS anviqo_plant_knowledge (knowledge_id TEXT PRIMARY KEY, organization_id TEXT NOT NULL, plant_id TEXT NOT NULL, document_id TEXT, record_type TEXT NOT NULL, external_id TEXT NOT NULL, name TEXT NOT NULL DEFAULT '', area TEXT NOT NULL DEFAULT '', service TEXT NOT NULL DEFAULT '', asset_type TEXT NOT NULL DEFAULT '', tag TEXT NOT NULL DEFAULT '', parent_id TEXT NOT NULL DEFAULT '', source TEXT NOT NULL DEFAULT '', metadata TEXT NOT NULL DEFAULT '{}', content TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL, UNIQUE(plant_id,external_id,source))")
            try: cur.execute("ALTER TABLE anviqo_plant_knowledge ADD COLUMN source TEXT NOT NULL DEFAULT ''")
            except Exception: pass
        else:
            cur.execute("CREATE TABLE IF NOT EXISTS anviqo_plant_knowledge (knowledge_id TEXT PRIMARY KEY, organization_id TEXT NOT NULL, plant_id TEXT NOT NULL, document_id TEXT, record_type TEXT NOT NULL, external_id TEXT NOT NULL, name TEXT NOT NULL DEFAULT '', area TEXT NOT NULL DEFAULT '', service TEXT NOT NULL DEFAULT '', asset_type TEXT NOT NULL DEFAULT '', tag TEXT NOT NULL DEFAULT '', parent_id TEXT NOT NULL DEFAULT '', source TEXT NOT NULL DEFAULT '', metadata JSONB NOT NULL DEFAULT '{}'::jsonb, content TEXT NOT NULL DEFAULT '', created_at TIMESTAMPTZ NOT NULL DEFAULT NOW())")
            cur.execute("ALTER TABLE anviqo_plant_knowledge ADD COLUMN IF NOT EXISTS source TEXT NOT NULL DEFAULT ''")
            cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS uq_anviqo_plant_knowledge_plant_external_source ON anviqo_plant_knowledge(plant_id,external_id,source)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_anviqo_plant_knowledge_plant ON anviqo_plant_knowledge(plant_id)")

def _insert_batch(plant_id,org_id,document_id,digest,records):
    p=_p(); rows=[]; errors=[]
    for raw in records:
        try:
            normalized=normalize_record(raw,str(raw.get("record_type") or "ASSET")); item=asdict(normalized) if is_dataclass(normalized) else dict(normalized)
            meta=dict(item.get("metadata") or {}); meta["import_version"]=VERSION; meta["document_sha256"]=digest; content=str(meta.pop("content","") or "")
            kid="know_"+hashlib.sha256((plant_id+document_id+str(item.get("external_id","")+str(item.get("source","")))).encode()).hexdigest()[:24]
            rows.append((kid,org_id,plant_id,document_id,item.get("record_type","ASSET"),str(item.get("external_id","")),item.get("name",""),item.get("area",""),item.get("service",""),item.get("asset_type",""),item.get("tag",""),item.get("parent_id",""),item.get("source",""),json.dumps(meta,default=str),content))
        except Exception as exc: errors.append(str(exc))
    if not rows:return 0,errors
    with store._connect() as conn:
        cur=conn.cursor()
        if store._is_sqlite(): cur.executemany(f"INSERT OR REPLACE INTO anviqo_plant_knowledge(knowledge_id,organization_id,plant_id,document_id,record_type,external_id,name,area,service,asset_type,tag,parent_id,source,metadata,content,created_at) VALUES({','.join([p]*15)},{p})",[r+(_now(),) for r in rows])
        else:
            placeholders=[p]*15; placeholders[13]=f"{p}::jsonb"; sql=f"INSERT INTO anviqo_plant_knowledge(knowledge_id,organization_id,plant_id,document_id,record_type,external_id,name,area,service,asset_type,tag,parent_id,source,metadata,content) VALUES({','.join(placeholders)}) ON CONFLICT(plant_id,external_id,source) DO UPDATE SET organization_id=EXCLUDED.organization_id,document_id=EXCLUDED.document_id,name=EXCLUDED.name,area=EXCLUDED.area,service=EXCLUDED.service,asset_type=EXCLUDED.asset_type,tag=EXCLUDED.tag,parent_id=EXCLUDED.parent_id,metadata=EXCLUDED.metadata,content=EXCLUDED.content"; cur.executemany(sql,rows)
    return len(rows),errors

def import_plant_data(plant_id,actor):
    if not actor.get("user_id") or actor.get("role") not in {"OWNER","ADMIN"}: raise PermissionError("OWNER/ADMIN access to this plant is required")
    p=_p(); _knowledge_schema()
    with store._connect() as conn:
        cur=conn.cursor(); cur.execute(f"SELECT 1 FROM anviqo_plants WHERE plant_id={p} AND organization_id={p} AND status='ACTIVE'",(plant_id,actor["organization_id"]))
        if not cur.fetchone(): raise PermissionError("Plant is not active in this organization")
        cur.execute(f"SELECT document_id,filename,sha256 FROM anviqo_plant_documents WHERE plant_id={p} AND organization_id={p} ORDER BY created_at",(plant_id,actor["organization_id"])); docs=cur.fetchall()
    total=documents=0; errors=[]
    for document_id,filename,digest in docs:
        try:
            with store._connect() as conn:
                cur=conn.cursor(); cur.execute(f"SELECT content FROM anviqo_plant_documents WHERE document_id={p} AND plant_id={p} AND organization_id={p}",(document_id,plant_id,actor["organization_id"])); row=cur.fetchone()
            if not row: raise RuntimeError("Document content not found")
            raw=bytes(row[0]); batch=[]
            for rec in _records(filename,raw):
                batch.append(rec)
                if len(batch)>=50:
                    added,batch_errors=_insert_batch(plant_id,actor["organization_id"],document_id,digest,batch); total+=added; errors.extend({"filename":filename,"error":e} for e in batch_errors); batch=[]
            if batch:
                added,batch_errors=_insert_batch(plant_id,actor["organization_id"],document_id,digest,batch); total+=added; errors.extend({"filename":filename,"error":e} for e in batch_errors)
            documents+=1; del raw
        except Exception as exc: errors.append({"filename":filename,"error":str(exc)})
    status="READY" if total>0 else "BLOCKED" if errors else "DATA_UPLOADED"
    with store._connect() as conn: conn.cursor().execute(f"UPDATE anviqo_plant_onboarding SET status={p},updated_at={p} WHERE plant_id={p} AND organization_id={p}",(status,_now(),plant_id,actor["organization_id"]))
    try: store.record_audit(actor,"IMPORT_PLANT_DATA","PLANT",plant_id,{"documents":documents,"records":total,"errors":len(errors),"version":VERSION})
    except Exception: pass
    return {"status":"OK" if total>0 else "NO_KNOWLEDGE_CREATED","plant_id":plant_id,"documents_processed":documents,"records_available_to_anvi":total,"errors":errors[:100],"error_count":len(errors),"import_version":VERSION,"mode":"DIRECT","message":("Imported knowledge successfully" if total>0 else ("Documents were found but produced zero normalized knowledge records. See errors for the exact document/parser/database reason." if errors else "Documents were found but the parser produced zero records. Check the uploaded file format/content.")),"safety":dict(SAFETY)}

def register(app):
    from flask import jsonify,session
    @app.post("/api/admin/onboarding/plant/<plant_id>/import")
    def import_route(plant_id):
        actor={"user_id":session.get("user_id",""),"organization_id":session.get("organization_id",""),"plant_id":session.get("plant_id",""),"role":session.get("role",""),"username":session.get("username","")}
        try:
            result=import_plant_data(plant_id,actor); return jsonify(result),200 if result["status"]=="OK" else 422
        except Exception as exc:return jsonify({"status":"ERROR","message":str(exc),"mode":"DIRECT","safety":dict(SAFETY)}),400
    @app.get("/api/admin/onboarding/plant/<plant_id>/knowledge-summary")
    def knowledge_summary(plant_id):
        actor={"user_id":session.get("user_id",""),"organization_id":session.get("organization_id",""),"plant_id":session.get("plant_id",""),"role":session.get("role",""),"username":session.get("username","")}
        if not actor.get("user_id") or actor.get("role") not in {"OWNER","ADMIN"}: return jsonify({"status":"FORBIDDEN"}),403
        p=_p()
        with store._connect() as conn:
            cur=conn.cursor(); cur.execute(f"SELECT record_type,COUNT(*) FROM anviqo_plant_knowledge WHERE plant_id={p} AND organization_id={p} GROUP BY record_type ORDER BY COUNT(*) DESC",(plant_id,actor["organization_id"])); groups=[{"record_type":r[0],"count":r[1]} for r in cur.fetchall()]; cur.execute(f"SELECT COUNT(*) FROM anviqo_plant_knowledge WHERE plant_id={p} AND organization_id={p}",(plant_id,actor["organization_id"])); total=cur.fetchone()[0]
        return jsonify({"status":"OK","plant_id":plant_id,"total_records":total,"by_type":groups,"mode":"DIRECT","safety":dict(SAFETY)})
    return app

__all__=["import_plant_data","register"]