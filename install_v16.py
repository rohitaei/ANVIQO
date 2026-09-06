from pathlib import Path
import shutil
from datetime import datetime

p = Path("pci_spares.py")

stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
backup = p.with_name(f"pci_spares_BEFORE_V1_6_{stamp}.py")
shutil.copy2(p, backup)

s = p.read_text(encoding="utf-8")

if "ANVIQO CRITICAL SPARES V1.6" in s:
    print("V1.6 already installed")
    print("Backup:", backup)
    raise SystemExit

patch = r'''

# ================================================================
# ANVIQO CRITICAL SPARES V1.6
# SAFE CONVERSATIONAL INVENTORY MUTATION
# ================================================================

import re as _v16_re
import openpyxl as _v16_openpyxl
from datetime import datetime as _v16_datetime
from pathlib import Path as _v16_Path

_V16_XLSX = _v16_Path(__file__).resolve().parent / \
    "database" / "spares" / "critical_spares.xlsx"

_V16_AUDIT = _V16_XLSX.parent / "spare_inventory_audit.log"


def _v16_action(q):
    q = str(q or "").lower()

    if any(x in q for x in (
        "add spare", "add spares",
        "new spare", "new spares",
        "receive spare", "receive spares",
        "received spare", "received spares",
        "increase stock", "increase spare"
    )):
        return "ADD"

    if any(x in q for x in (
        "use spare", "use spares",
        "used spare", "used spares",
        "consume spare", "consumed spare",
        "issue spare", "issued spare",
        "remove spare", "remove spares",
        "decrease stock"
    )):
        return "USE"

    return None


def _v16_quantity(q):
    q = str(q or "")

    patterns = [
        r'\b(\d+)\s*(?:nos?|numbers?|pcs?|pieces?|qty|quantity)\b',
        r'\b(?:add|receive|received|increase|put|use|used|consume|issue|issued|remove)\s+(\d+)\b'
    ]

    for pattern in patterns:
        m = _v16_re.search(pattern, q, _v16_re.I)
        if m:
            n = int(m.group(1))
            if n > 0:
                return n

    return None


def _v16_identifier(q):
    q = str(q or "")

    # Exact equipment tags first: PT-303, MCV-205 etc.
    m = _v16_re.search(
        r'\b([A-Z]{2,8}-\d{2,5})\b',
        q,
        _v16_re.I
    )

    if m:
        return m.group(1).upper()

    # Untagged category commands such as "receive 5 SOV spares"
    for family in (
        "SOV", "FSV", "PCV", "FCV", "MCV",
        "PT", "LT", "FT", "RTD"
    ):
        if family in q.upper():
            return family

    return None


def _v16_find(identifier):
    rows = _v14_load_spares()
    ident = str(identifier).upper().replace("-", "")

    matches = []

    for r in rows:
        tag = str(r.get("tag") or "").upper().replace("-", "")
        instrument = str(r.get("instrument") or "").upper().replace("-", "")

        if ident == tag or ident == instrument:
            matches.append(r)

    return matches


def execute_spare_mutation(question, confirmed=False):

    # SAFETY GATE
    if not confirmed:
        return {
            "ok": False,
            "executed": False,
            "human_decision_required": True,
            "plc_write": False,
            "scada_control": False,
            "message": "Confirmation required. No inventory changed."
        }

    action = _v16_action(question)
    quantity = _v16_quantity(question)
    identifier = _v16_identifier(question)

    if not action:
        return {
            "ok": False,
            "executed": False,
            "error": "No supported spare action detected."
        }

    if not quantity:
        return {
            "ok": False,
            "executed": False,
            "error": "Positive quantity required."
        }

    if not identifier:
        return {
            "ok": False,
            "executed": False,
            "error": "Spare tag/item identifier required."
        }

    matches = _v16_find(identifier)

    if len(matches) == 0:
        return {
            "ok": False,
            "executed": False,
            "error": f"No exact spare record found for {identifier}."
        }

    if len(matches) > 1:
        return {
            "ok": False,
            "executed": False,
            "ambiguous": True,
            "matches": [
                {
                    "tag": r.get("tag"),
                    "instrument": r.get("instrument"),
                    "sheet": r.get("sheet"),
                    "row": r.get("row"),
                    "available": r.get("qty_available")
                }
                for r in matches
            ],
            "error": "Multiple matching spare records. No inventory changed."
        }

    r = matches[0]
    sheet = r["sheet"]
    row = int(r["row"])
    before = float(r.get("qty_available") or 0)

    if action == "USE":
        if quantity > before:
            return {
                "ok": False,
                "executed": False,
                "error": (
                    f"Insufficient stock. Available {before:g}, "
                    f"requested {quantity}. No inventory changed."
                )
            }
        after = before - quantity
    else:
        after = before + quantity

    wb = _v16_openpyxl.load_workbook(_V16_XLSX)
    ws = wb[sheet]

    # Find the existing available quantity by matching the value
    # identified by the robust reader.
    target_col = None

    for col in range(1, ws.max_column + 1):
        value = ws.cell(row, col).value

        try:
            if value is not None and float(value) == before:
                target_col = col
                break
        except Exception:
            pass

    if target_col is None:
        wb.close()
        return {
            "ok": False,
            "executed": False,
            "error": "Could not safely identify stock cell. No inventory changed."
        }

    ws.cell(row, target_col).value = int(after) if after.is_integer() else after

    wb.save(_V16_XLSX)
    wb.close()

    # Verify from disk
    check = _v16_openpyxl.load_workbook(
        _V16_XLSX,
        data_only=True,
        read_only=True
    )

    verified = check[sheet].cell(row, target_col).value
    check.close()

    if float(verified) != float(after):
        return {
            "ok": False,
            "executed": False,
            "error": "Excel verification failed."
        }

    timestamp = _v16_datetime.now().isoformat(timespec="seconds")

    with open(_V16_AUDIT, "a", encoding="utf-8") as f:
        f.write(
            f"{timestamp} | {action} | "
            f"{identifier} | Sheet={sheet} | Row={row} | "
            f"Before={before:g} | Quantity={quantity} | "
            f"After={after:g}\n"
        )

    return {
        "ok": True,
        "executed": True,
        "action": action,
        "tag": identifier,
        "quantity": quantity,
        "before": before,
        "after": after,
        "sheet": sheet,
        "row": row,
        "verified": True,
        "audit_log": str(_V16_AUDIT),
        "plc_write": False,
        "scada_control": False,
        "human_decision_required": True,
        "answer": (
            f"ANVIQO inventory updated successfully.\n"
            f"{identifier}: {before:g} -> {after:g}\n"
            f"Action: {action}; Quantity: {quantity}\n"
            f"Excel verified: YES\n"
            f"PLC write: FALSE\n"
            f"SCADA control: FALSE"
        )
    }


def answer_spare_management_v16(question, confirmed=False):

    if _v16_action(question):
        return execute_spare_mutation(
            question,
            confirmed=confirmed
        )

    return answer_spare_query(question)

'''

p.write_text(s + patch, encoding="utf-8")

print("==============================================================")
print("ANVIQO CRITICAL SPARES V1.6 INSTALLED")
print("==============================================================")
print("Backup:", backup)
print("PLC write       : FALSE")
print("SCADA control   : FALSE")
print("Human decision  : REQUIRED")
print("Excel mutation  : ENABLED AFTER CONFIRMATION")
