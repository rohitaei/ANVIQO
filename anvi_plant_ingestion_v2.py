"""Durable, low-memory tenant plant ingestion adapter.

This is data ingestion only. Existing V5 intelligence and PLC/SCADA safety
boundaries are untouched. CHANGE DATA, NOT CODE.
"""
from __future__ import annotations

import hashlib, io, json, threading
from datetime import datetime, timezone

import anvi_tenant_store as store
from universal_onboarding import build_onboarding_package, SAFETY

JOB_TABLE = "anviqo_plant_ingestion_jobs"
LOCK = threading.Lock()
STALE_SECONDS = 15 * 60


def _p(): return store._placeholder()
def _now(): return datetime.now(timezone.utc).isoformat()

def _actor():
    from flask import session
    return {"user_id":session.get("user_id",""),"organization_id":session.get("organization_id",""),"plant_id":session.get("plant_id",""),"role":session.get("role","")}

def _allowed(plant_id, a):
    if not a.get("user_id") or not a.get("organization_id") or a.get("role") not in {"OWNER","ADMIN"}: return False
    with store._connect() as c:
        q=c.cursor(); q.execute(f"SELECT 1 FROM anviqo_plants WHERE plant_id={_p()} AND organization_id={_p()} AND status='ACTIVE' LIMIT 1",(plant_id,a["organization_id"]))
        return q.fetchone() is not None

def init_schema():
    if not store.enabled(): return
    with store._connect() as c:
        q=c.cursor()
        if store._is_sqlite():
            q.execute("CREATE TABLE IF NOT EXISTS anviqo_plant_ingestion_jobs (plant_id TEXT PRIMARY KEY, organization_id TEXT NOT NULL, status TEXT NOT NULL, message TEXT NOT NULL DEFAULT '', documents_processed INTEGER NOT NULL DEFAULT 0, records_indexed INTEGER NOT NULL DEFAULT 0, errors TEXT NOT NULL DEFAULT '[]', started_at TEXT, finished_at TEXT, updated_at TEXT NOT NULL)")
        else:
            q.execute("CREATE TABLE IF NOT EXISTS anviqo_plant_ingestion_jobs (plant_id TEXT PRIMARY KEY REFERENCES anviqo_plants(plant_id), organization_id TEXT NOT NULL, status TEXT NOT NULL, message TEXT NOT NULL DEFAULT '', documents_processed INTEGER NOT NULL DEFAULT 0, records_indexed INTEGER NOT NULL DEFAULT 0, errors JSONB NOT NULL DEFAULT '[]'::jsonb, started_at TIMESTAMPTZ, finished_at TIMESTAMPTZ, updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW())")

def _save(plant, org, status, message, docs=0, records=0, errors=None, started=None, finished=None):
    init_schema(); now=_now(); errors=errors or []
    with store._connect() as c:
        q=c.cursor(); vals=(plant,org,status,message,docs,records,json.dumps(errors),started,finished,now)
        if store._is_sqlite():
            q.execute(f"INSERT OR REPLACE INTO {JOB_TABLE}(plant_id,organization_id,status,message,documents_processed,records_indexed,errors,started_at,finished_at,updated_at) VALUES({','.join([_p()]*10)})",vals)
        else:
            q.execute(f"INSERT INTO {JOB_TABLE}(plant_id,organization_id,status,message,documents_processed,records_indexed,errors,started_at,finished_at,updated_at) VALUES({','.join([_p()]*10)}) ON CONFLICT(plant_id) DO UPDATE SET status=EXCLUDED.status,message=EXCLUDED.message,documents_processed=EXCLUDED.documents_processed,records_indexed=EXCLUDED.records_indexed,errors=EXCLUDED.errors,started_at=EXCLUDED.started_at,finished_at=EXCLUDED.finished_at,updated_at=EXCLUDED.updated_at",vals)

def _job(plant,org):
    init_schema()
    with store._connect() as c:
        q=c.cursor(); q.execute(f"SELECT status,message,documents_processed,records_indexed,errors,started_at,finished_at,updated_at FROM {JOB_TABLE} WHERE plant_id={_p()} AND organization_id={_p()}",(plant,org)); r=q.fetchone()
    if not r: return None
    try: errors=json.loads(r[4]) if isinstance(r[4],str) else (r[4] or [])
    except Exception: errors=[]
    status=r[0]; message=r[1]
    if status=="RUNNING" and r[7]:
        try:
            age=(datetime.now(timezone.utc)-datetime.fromisoformat(str(r[7]).replace("Z","+00:00"))).total_seconds()
            if age>STALE_SECONDS:
                status="STALE"; message="Previous ingestion worker stopped before completion. It is safe to start ingestion again."
        except Exception: pass
    return {"status":status,"job_status":status,"plant_id":plant,"message":message,"documents_processed":r[2],"records_indexed":r[3],"errors":errors,"started_at":r[5],"finished_at":r[6],"updated_at":r[7],"safety":dict(SAFETY)}

def _xlsx_batches(raw, batch_size=500):
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    try:
        batch=[]
        for ws in wb.worksheets:
            rows=ws.iter_rows(values_only=True)
            try: first=next(rows)
            except StopIteration: continue
            headers=[str(v).strip() if v is not None and str(v).strip() else f"column_{i+1}" for i,v in enumerate(first)]
            for idx,row in enumerate(rows,start=2):
                vals=list(row)
                if not any(v is not None and str(v).strip() for v in vals): continue
                rec={headers[i]:vals[i] for i in range(min(len(headers),len(vals)))}
                if not rec: continue
                rec["external_id"]=str(rec.get("external_id") or rec.get("id") or rec.get("tag") or rec.get("asset_id") or rec.get("PLC TAG") or rec.get("PLC TAG NAME") or rec.get("TAG NAME") or hashlib.sha256((ws.title+str(idx)+json.dumps(rec,sort_keys=True,default=str)).encode()).hexdigest()[:20])
                rec["source"]=ws.title
                rec["record_type"]=str(rec.get("record_type") or rec.get("entity_type") or rec.get("type") or "ASSET")
                batch.append(rec)
                if len(batch)>=batch_size:
                    yield batch; batch=[]
        if batch: yield batch
    finally: wb.close()

def _batches(filename, raw):
    ext=filename.lower().rsplit('.',1)[-1] if '.' in filename else ''
    if ext in {'xlsx','xls'}: return _xlsx_batches(raw)
    from anvi_plant_ingestion_runtime import _records
    records=_records(filename,raw)
    return (records[i:i+500] for i in range(0,len(records),500))

def _insert_rows(rows):
    if not rows: return
    p=_p()
    with store._connect() as c:
        q=c.cursor()
        if store._is_sqlite():
            q.executemany(f"INSERT OR REPLACE INTO anviqo_plant_knowledge(knowledge_id,organization_id,plant_id,document_id,record_type,external_id,name,area,service,asset_type,tag,parent_id,source,metadata,content,created_at) VALUES({','.join([p]*15)},{p})",[r+(_now(),) for r in rows])
        else:
            q.executemany(f"INSERT INTO anviqo_plant_knowledge(knowledge_id,organization_id,plant_id,document_id,record_type,external_id,name,area,service,asset_type,tag,parent_id,source,metadata,content) VALUES({','.join([p]*15)}) ON CONFLICT(plant_id,external_id,source) DO UPDATE SET name=EXCLUDED.name,area=EXCLUDED.area,service=EXCLUDED.service,asset_type=EXCLUDED.asset_type,tag=EXCLUDED.tag,metadata=EXCLUDED.metadata,content=EXCLUDED.content",rows)

def ingest(plant_id,a):
    if not _allowed(plant_id,a): raise PermissionError("OWNER/ADMIN access to this plant is required")
    init_schema(); p=_p(); started=_now(); _save(plant_id,a["organization_id"],"RUNNING","Universal ingestion is processing plant documents.",started=started)
    with store._connect() as c:
        q=c.cursor(); q.execute(f"SELECT p.name,o.industry FROM anviqo_plants p LEFT JOIN anviqo_plant_onboarding o ON o.plant_id=p.plant_id WHERE p.plant_id={p} AND p.organization_id={p}",(plant_id,a["organization_id"])); plant=q.fetchone()
        q.execute(f"SELECT document_id,filename,content,sha256 FROM anviqo_plant_documents WHERE plant_id={p} AND organization_id={p} ORDER BY created_at",(plant_id,a["organization_id"])); docs=q.fetchall()
    if not plant: raise ValueError("Plant not found")
    total=documents=0; errors=[]
    for document_id,filename,raw,digest in docs:
        try:
            identity={"organization_id":a["organization_id"],"plant_id":plant_id,"name":plant[0],"industry":plant[1] or ""}
            for batch in _batches(filename,bytes(raw)):
                package=build_onboarding_package(identity,batch)
                rows=[]
                for rec in package.get("records",[]):
                    kid="know_"+hashlib.sha256((plant_id+document_id+rec["external_id"]+rec.get("source","")).encode()).hexdigest()[:24]
                    meta=dict(rec.get("metadata") or {}); meta.update({"ingestion_version":"ANVIQO-PLANT-INGESTION-V2","document_sha256":digest}); content=str(meta.pop("content","") or "")
                    rows.append((kid,a["organization_id"],plant_id,document_id,rec["record_type"],rec["external_id"],rec.get("name",""),rec.get("area",""),rec.get("service",""),rec.get("asset_type",""),rec.get("tag",""),rec.get("parent_id",""),rec.get("source",filename),json.dumps(meta),content))
                _insert_rows(rows); total+=len(rows)
            documents+=1; _save(plant_id,a["organization_id"],"RUNNING",f"Processed {documents} document(s).",documents,total,errors,started)
        except Exception as exc: errors.append({"filename":filename,"error":str(exc)})
    status="READY" if documents and not errors else "BLOCKED" if errors and not total else "DATA_UPLOADED"
    _save(plant_id,a["organization_id"],status,"Universal ingestion completed." if status=="READY" else "Universal ingestion completed with errors.",documents,total,errors,started,_now())
    with store._connect() as c:
        q=c.cursor(); q.execute(f"UPDATE anviqo_plant_onboarding SET status={p},updated_at={p} WHERE plant_id={p}",(status,_now(),plant_id))
    return {"status":"OK","plant_id":plant_id,"documents_processed":documents,"records_indexed":total,"errors":errors,"ingestion_version":"ANVIQO-PLANT-INGESTION-V2","safety":dict(SAFETY)}

def _worker(plant,a):
    try: result=ingest(plant,a); result["job_status"]="COMPLETED"
    except Exception as exc:
        result={"status":"ERROR","plant_id":plant,"message":str(exc),"job_status":"FAILED","safety":dict(SAFETY)}
        _save(plant,a["organization_id"],"FAILED",str(exc),errors=[{"error":str(exc)}])

def register(app):
    from flask import jsonify
    @app.post('/api/admin/onboarding/plant/<plant_id>/ingest')
    def start(plant_id):
        a=_actor()
        if not _allowed(plant_id,a): return jsonify({"status":"FORBIDDEN"}),403
        j=_job(plant_id,a["organization_id"])
        if j and j["status"]=="RUNNING": return jsonify(j),202
        _save(plant_id,a["organization_id"],"RUNNING","Universal ingestion started.",started=_now())
        with store._connect() as c:
            q=c.cursor(); q.execute(f"UPDATE anviqo_plant_onboarding SET status='INDEXING',updated_at={_p()} WHERE plant_id={_p()}",(_now(),plant_id))
        threading.Thread(target=_worker,args=(plant_id,a),daemon=True,name=f'anvi-ingest-{plant_id}').start()
        return jsonify({"status":"STARTED","job_status":"RUNNING","plant_id":plant_id,"message":"Universal ingestion started. Monitor Onboarding Readiness for completion.","safety":dict(SAFETY)}),202
    @app.get('/api/admin/onboarding/plant/<plant_id>/ingest-status')
    def ingestion_status_v2(plant_id):
        a=_actor()
        if not _allowed(plant_id,a): return jsonify({"status":"FORBIDDEN"}),403
        return jsonify(_job(plant_id,a["organization_id"]) or {"status":"IDLE","job_status":"IDLE","plant_id":plant_id,"safety":dict(SAFETY)})
    return app

__all__=['register','init_schema','ingest']
